"""
markpdf.pptx_converter - Native PowerPoint (.pptx) conversion to Markdown and text with offline diagram OCR.
"""

import io
import os
import sys
from typing import Optional, List, Any

try:
    import pptx
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

from .ocr import is_ocr_available, ocr_image_pil
from .tables import format_markdown_table


def is_pptx_supported() -> bool:
    """Check if python-pptx is installed."""
    return HAS_PPTX


def extract_slide_elements(
    slide: Any,
    slide_num: int,
    output_format: str = "md",
    ocr_graphics: bool = True,
    lang: str = "eng",
    label_graphics: bool = True,
) -> str:
    """
    Extract text, tables, and diagram OCR from a single slide in reading order.
    """
    from PIL import Image

    is_md = output_format == "md"
    title_text = ""
    slide_elements = []

    # Get shapes sorted by top, then left coordinate
    shapes = list(slide.shapes)
    shapes.sort(key=lambda s: (s.top if hasattr(s, "top") and s.top is not None else 0,
                              s.left if hasattr(s, "left") and s.left is not None else 0))

    title_shape = None
    if hasattr(slide.shapes, "title") and slide.shapes.title:
        title_shape = slide.shapes.title
        if title_shape.has_text_frame:
            title_text = title_shape.text_frame.text.strip()

    for shape in shapes:
        # Skip title shape since we handle it at the top of the slide
        if title_shape and shape == title_shape:
            continue

        # 1. Table
        if shape.has_table:
            rows = []
            for row in shape.table.rows:
                rows.append([cell.text.strip() for cell in row.cells])
            md_tab = format_markdown_table(rows)
            if md_tab:
                slide_elements.append(md_tab)
            continue

        # 2. Text Frame
        if shape.has_text_frame:
            tf_lines = []
            for p in shape.text_frame.paragraphs:
                p_text = p.text.strip()
                if not p_text:
                    continue
                # If bullet point level > 0
                if is_md and p.level > 0:
                    indent = "  " * p.level
                    tf_lines.append(f"{indent}- {p_text}")
                elif is_md and hasattr(shape, "placeholder_format") and shape.placeholder_format.type == 2:  # body placeholder
                    tf_lines.append(f"- {p_text}")
                else:
                    tf_lines.append(p_text)

            if tf_lines:
                slide_elements.append("\n".join(tf_lines))
            continue

        # 3. Picture / Diagram OCR
        if ocr_graphics and is_ocr_available():
            image_blob = None
            try:
                if hasattr(shape, "image") and shape.image:
                    image_blob = shape.image.blob
            except Exception:
                image_blob = None

            if image_blob:
                try:
                    img_obj = Image.open(io.BytesIO(image_blob))
                    # Only OCR pictures of reasonable size
                    if img_obj.width >= 40 and img_obj.height >= 40:
                        txt = ocr_image_pil(img_obj, lang=lang, upscale=True)
                        if txt:
                            if is_md:
                                content = f"> **[Diagram/Graphic OCR]**\n" + "\n".join(f"> {l}" for l in txt.split("\n")) if label_graphics else txt
                            else:
                                content = f"[Diagram/Graphic OCR]:\n{txt}" if label_graphics else txt
                            slide_elements.append(content)
                except Exception:
                    pass

    # Build slide text
    slide_lines = []

    # Slide Header
    if is_md:
        if title_text:
            slide_lines.append(f"## Slide {slide_num}: {title_text}")
        else:
            slide_lines.append(f"## Slide {slide_num}")
    else:
        if title_text:
            slide_lines.append(f"--- Slide {slide_num}: {title_text} ---")
        else:
            slide_lines.append(f"--- Slide {slide_num} ---")

    if slide_elements:
        slide_lines.append("")
        slide_lines.append("\n\n".join(slide_elements))

    # 4. Speaker Notes
    try:
        if hasattr(slide, "has_notes_slide") and slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            notes_text = notes_tf.text.strip()
            if notes_text:
                slide_lines.append("")
                if is_md:
                    notes_block = "> **Speaker Notes:**\n" + "\n".join(f"> {l}" for l in notes_text.split("\n"))
                    slide_lines.append(notes_block)
                else:
                    slide_lines.append(f"Speaker Notes:\n{notes_text}")
    except Exception:
        pass

    return "\n".join(slide_lines).strip()


def convert_pptx(
    pptx_path: str,
    output_path: Optional[str] = None,
    output_format: str = "md",
    ocr_graphics: bool = True,
    lang: str = "eng",
    label_graphics: bool = True,
    verbose: bool = True,
    add_page_markers: bool = True,
) -> Optional[str]:
    """
    Convert a PowerPoint (.pptx) presentation to Markdown or plain text.
    """
    if not os.path.exists(pptx_path):
        print(f"Error: Presentation file not found at '{pptx_path}'", file=sys.stderr)
        return None

    if not HAS_PPTX:
        print("Error: 'python-pptx' is not installed. Run: pip install python-pptx", file=sys.stderr)
        return None

    try:
        prs = pptx.Presentation(pptx_path)
        total_slides = len(prs.slides)

        if verbose:
            ocr_str = ", diagram OCR: on" if ocr_graphics and is_ocr_available() else ""
            print(f"Processing presentation '{pptx_path}' ({total_slides} slide{'s' if total_slides != 1 else ''}, format: {output_format}{ocr_str})...")

        slides_text = []
        for idx, slide in enumerate(prs.slides):
            slide_num = idx + 1
            s_text = extract_slide_elements(
                slide=slide,
                slide_num=slide_num,
                output_format=output_format,
                ocr_graphics=ocr_graphics,
                lang=lang,
                label_graphics=label_graphics,
            )
            slides_text.append(s_text)

        is_md = output_format == "md"
        separator = "\n\n---\n\n" if is_md else "\n\n"
        combined_text = separator.join(slides_text).strip() + "\n"

        ext = ".md" if is_md else ".txt"
        if output_path:
            if os.path.isdir(output_path):
                base_name = os.path.splitext(os.path.basename(pptx_path))[0]
                final_output_path = os.path.join(output_path, f"{base_name}{ext}")
            else:
                final_output_path = output_path
        else:
            base_name = os.path.splitext(pptx_path)[0]
            final_output_path = f"{base_name}{ext}"

        with open(final_output_path, "w", encoding="utf-8") as out_f:
            out_f.write(combined_text)

        print(f"Successfully converted '{pptx_path}' to '{final_output_path}'")
        return final_output_path

    except Exception as e:
        print(f"An error occurred during presentation conversion: {e}", file=sys.stderr)
        return None
