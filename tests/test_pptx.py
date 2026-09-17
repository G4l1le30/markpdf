"""
Unit tests for PowerPoint (.pptx) conversion in markpdf.
"""

import io
import os
import tempfile
import unittest

from markpdf.pptx_converter import is_pptx_supported, convert_pptx
from markpdf.converter import convert_pdf_to_text


class TestPPTXConverter(unittest.TestCase):
    def setUp(self):
        if not is_pptx_supported():
            self.skipTest("python-pptx not installed")

    def test_pptx_conversion_complete(self):
        import pptx
        from pptx.util import Inches
        from PIL import Image, ImageDraw

        with tempfile.TemporaryDirectory() as tmpdir:
            pptx_path = os.path.join(tmpdir, "sample.pptx")
            prs = pptx.Presentation()

            # Slide 1: Title, Bullets, Diagram Image, and Speaker Notes
            slide_layout = prs.slide_layouts[1]
            slide1 = prs.slides.add_slide(slide_layout)
            slide1.shapes.title.text = "System Architecture"
            tf = slide1.shapes.placeholders[1].text_frame
            tf.text = "Overview of the deployment"
            p = tf.add_paragraph()
            p.text = "High availability clustering"
            p.level = 1

            # Diagram Image on Slide 1
            img = Image.new("RGB", (300, 100), color=(255, 255, 255))
            d = ImageDraw.Draw(img)
            d.text((20, 40), "Load Balancer Node", fill=(0, 0, 0))
            img_buf = io.BytesIO()
            img.save(img_buf, format="PNG")
            img_buf.seek(0)
            slide1.shapes.add_picture(img_buf, Inches(1), Inches(4), width=Inches(4))

            # Speaker Notes on Slide 1
            slide1.notes_slide.notes_text_frame.text = "Highlight zero-downtime failover."

            # Slide 2: Table
            blank_layout = prs.slide_layouts[6]
            slide2 = prs.slides.add_slide(blank_layout)
            table_shape = slide2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(5), Inches(1.5))
            t = table_shape.table
            t.cell(0, 0).text = "Service"
            t.cell(0, 1).text = "Port"
            t.cell(1, 0).text = "Wazuh API"
            t.cell(1, 1).text = "55000"

            prs.save(pptx_path)

            # Convert to Markdown via generic entrypoint
            md_out = convert_pdf_to_text(pptx_path, output_format="md", verbose=False)
            self.assertTrue(os.path.exists(md_out))
            self.assertTrue(md_out.endswith(".md"))

            with open(md_out, "r", encoding="utf-8") as f:
                content = f.read()

            # Verify slide titles & content
            self.assertIn("## Slide 1: System Architecture", content)
            self.assertIn("Overview of the deployment", content)
            self.assertIn("  - High availability clustering", content)

            # Verify speaker notes
            self.assertIn("> **Speaker Notes:**", content)
            self.assertIn("Highlight zero-downtime failover.", content)

            # Verify table
            self.assertIn("| Service | Port |", content)
            self.assertIn("| Wazuh API | 55000 |", content)

            # Verify diagram OCR
            self.assertIn("> **[Diagram/Graphic OCR]**", content)


if __name__ == "__main__":
    unittest.main()
